from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_period_data_to_excel(
    parsed_data: Dict, output_path: str = "report.xlsx"
) -> str:
    """Форматує зібрані дані витрат у Excel-таблицю та зберігає їх у файл формату .xlsx.

    Створює новий документ Excel за допомогою openpyxl, записує загальні
    витрати за вказаний період та окремо витрати за кожним пошуковим шаблоном.
    Опціонально записує звіт по витратах поза пошуковими шаблонами.
    Застосовує грошове форматування числових комірок (`# ##0.00" грн."`),
    виділяє заголовки жирним шрифтом та автоматично розширює ширину колонок
    відповідно до довжини тексту.

    Args:
        parsed_data (Dict): Словник із результатами парсингу. Очікує наступну структуру:
            - "start_date" (str): Початкова дата періоду.
            - "end_date" (str): Кінцева дата періоду.
            - "total_spent" (float): Загальна сума витрат.
            - "template_results" (Dict[str, float]): Словник з результатами
              витрат по кожному текстовому шаблону.
            - "other_spent" (float, optional): Витрати поза шаблонами.
        output_path (Union[str, Path], optional): Шлях до файлу для збереження звіту.
            За замовчуванням "report.xlsx".

    Returns:
        str: Абсолютний шлях до створеного та збереженого файлу Excel.

    Raises:
        PermissionError: Якщо файл за вказаним шляхом відкритий в іншій програмі
            або відсутні права на запис у директорію.
        OSError: Якщо виникла системна помилка під час збереження файлу на диск.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Звіт TurboSMS"

    bold_font = Font(bold=True)
    currency_format = '# ##0.00" грн."'

    start_date = parsed_data.get("start_date", "")
    end_date = parsed_data.get("end_date", "")
    total_spent = parsed_data.get("total_spent", 0.0)
    template_results = parsed_data.get("template_results", {})

    # 1. Загальні витрати за період
    ws["A1"] = f"Всього витрачено за період {start_date} — {end_date}:"
    ws["A1"].font = bold_font
    ws["B1"] = total_spent
    ws["B1"].number_format = currency_format
    ws["B1"].font = bold_font

    # 2. Заголовок розділу шаблонів
    ws["A2"] = "Витрачено на повідомлення зі змістом:"
    ws["A2"].font = bold_font
    ws.merge_cells("A2:B2")

    # 3. Шаблони та їх суми
    current_row = 3
    for template_text, amount in template_results.items():
        ws[f"A{current_row}"] = f'"{template_text}":'
        ws[f"B{current_row}"] = amount
        ws[f"B{current_row}"].number_format = currency_format
        current_row += 1

    # 4. Витрати поза шаблонами (якщо розраховані)
    if "other_spent" in parsed_data:
        ws[f"A{current_row}"] = "Інші повідомлення:"
        ws[f"A{current_row}"].font = bold_font
        ws[f"B{current_row}"] = parsed_data["other_spent"]
        ws[f"B{current_row}"].number_format = currency_format
        ws[f"B{current_row}"].font = bold_font
        current_row += 1

    # 5. Автоматичне розширення колонок
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max(max_len + 3, 20)

    save_path = Path(output_path)
    wb.save(save_path)

    return str(save_path.resolve())


